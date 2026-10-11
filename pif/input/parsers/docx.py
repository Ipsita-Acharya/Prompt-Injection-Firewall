"""DOCX parser (Phase 2): paragraphs, tables, embedded media OCR, XML parts
and hidden-content forensics — runs flagged hidden (vanish), sub-2pt, or
white-on-white are extracted as their own segments, as are headers and
footers, which ordinary extraction never reaches."""

import io
from xml.etree import ElementTree

from pif.input.base import InspectionError, Parser, check_zip
from pif.input.ocr import faint_text_ocr, ocr_image
from pif.observability import content_preview, logger

_MAX_HIDDEN_CHARS = 20_000
_WHITE = (0xFF, 0xFF, 0xFF)


def _run_is_hidden(run) -> str | None:
    """Return the hidden-channel name when the run is invisible to a reader:
    the vanish flag, a sub-2pt font, or a white fill (the DOCX equivalents of
    the CSS hidden-content tricks)."""
    font = run.font
    if font.hidden:
        return "vanish"
    try:
        if font.size is not None and 0 < font.size.pt < 2.0:
            return "tiny"
    except (TypeError, ValueError):
        pass
    try:
        rgb = font.color.rgb
        if rgb is not None and (rgb[0], rgb[1], rgb[2]) == _WHITE:
            return "white"
    except (TypeError, ValueError):
        pass
    return None


def _hidden_run_segments(document) -> list[tuple[str, str]]:
    segments: list[tuple[str, str]] = []
    for paragraph in document.paragraphs:
        if len(segments) >= 48:
            break
        hidden_text = " ".join(
            run.text for run in paragraph.runs if run.text.strip() and _run_is_hidden(run)
        ).strip()
        if hidden_text:
            logger.info("docx.hidden_text_captured", extra={"event_fields": dict(channel="run_flags", text=content_preview(hidden_text))})
            segments.append((f"docx_hidden:{len(segments) + 1}", hidden_text[:_MAX_HIDDEN_CHARS]))
    return segments


def _header_footer_segments(document) -> list[tuple[str, str]]:
    """Headers/footers are content the model reads in page renderings but
    paragraph-only extraction never captures."""
    segments: list[tuple[str, str]] = []
    for index, section in enumerate(document.sections):
        for kind, container in (("header", section.header), ("footer", section.footer)):
            text = " ".join(
                paragraph.text for paragraph in container.paragraphs if paragraph.text.strip()
            ).strip()
            if text:
                segments.append((f"{kind}:{index + 1}", text[:_MAX_HIDDEN_CHARS]))
    return segments


class DocxParser(Parser):
    format_name = "docx"

    def parse(self, data: bytes, name: str, declared: str, settings, budget: dict, dispatcher, depth: int = 0):
        from docx import Document

        document = Document(io.BytesIO(data))
        logger.debug("docx.opened", extra={"event_fields": dict(paragraph_count=len(document.paragraphs), table_count=len(document.tables))})
        result = []
        for index, paragraph in enumerate(document.paragraphs):
            logger.info("docx.paragraph_captured", extra={"event_fields": dict(paragraph_index=index, text=content_preview(paragraph.text))})
            result.append((f"paragraph:{index}", paragraph.text))
        result.extend(_hidden_run_segments(document))
        result.extend(_header_footer_segments(document))
        for index, table in enumerate(document.tables):
            table_text = "\n".join(" ".join(cell.text for cell in row.cells) for row in table.rows)
            logger.info("docx.table_captured", extra={"event_fields": dict(table_index=index, text=content_preview(table_text))})
            result.append((f"table:{index}", table_text))
        with check_zip(data, settings) as archive:
            for entry in archive.namelist():
                if entry.startswith("word/media/"):
                    if not settings.ocr_enabled:
                        raise InspectionError("docx_images_require_ocr")
                    image_data = archive.read(entry)
                    result.append((f"ocr:{entry}", ocr_image(image_data, settings)))
                    try:
                        faint = faint_text_ocr(image_data, exclude=set(result[-1][1].split()), settings=settings)
                    except InspectionError:
                        faint = ""
                    if faint:
                        logger.info("docx.hidden_text_captured", extra={"event_fields": dict(channel="faint_ocr", entry=entry, text=content_preview(faint))})
                        result.append((f"ocr_faint:{entry}", faint[:_MAX_HIDDEN_CHARS]))
                if entry.endswith(".xml") and entry.startswith(("word/", "docProps/")):
                    xml = archive.read(entry)
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise InspectionError("xml_entities_unsupported")
                    root = ElementTree.fromstring(xml)
                    text = " ".join(root.itertext())
                    attributes = " ".join(
                        value
                        for element in root.iter()
                        for key, value in element.attrib.items()
                        if key.rsplit("}", 1)[-1]
                        in {"descr", "title", "tooltip", "author", "name", "alt"}
                    )
                    logger.info("docx.xml_captured", extra={"event_fields": dict(entry=entry, text=content_preview(text + " " + attributes))})
                    result.append((f"xml:{entry}", text + " " + attributes))
        return result
